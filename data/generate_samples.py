"""
Generates realistic sample resumes (PDF, DOCX, TXT) and job descriptions
across Data Science, DevOps, and Backend Engineering for realistic testing and evaluation.
"""

import os
from fpdf import FPDF
import docx

DATA_DIR = os.path.dirname(os.path.abspath(__file__))
SAMPLES_DIR = os.path.join(DATA_DIR, "samples")
RESUMES_DIR = os.path.join(SAMPLES_DIR, "resumes")
JDS_DIR = os.path.join(SAMPLES_DIR, "job_descriptions")

os.makedirs(RESUMES_DIR, exist_ok=True)
os.makedirs(JDS_DIR, exist_ok=True)

# 1. Data Scientist Resume Content
DATA_SCIENTIST_RESUME = """
ALEX CHEN
San Francisco, CA | alex.chen@email.com | github.com/alexchen | linkedin.com/in/alexchen

PROFESSIONAL SUMMARY
Senior Machine Learning Engineer & Data Scientist with 5+ years of experience building, deploying, and scaling predictive models and Generative AI applications. Proficient in PyTorch, Scikit-Learn, and NLP pipelines using Transformers, BERT, and LangChain. Proven track record of architecting RAG systems and optimizing vector search across high-throughput production environments on AWS.

TECHNICAL SKILLS
• Languages: Python, SQL, R, Bash
• AI & Machine Learning: Machine Learning, Deep Learning, NLP, LLM, RAG, Transformers, BERT, Computer Vision, Feature Engineering, Supervised Learning, Unsupervised Learning, Fine-Tuning
• Frameworks & Libraries: PyTorch, Scikit-Learn, Hugging Face, LangChain, Pandas, NumPy, SciPy, OpenCV, FastAPI
• Databases & Vector Stores: PostgreSQL, MongoDB, Redis, Pinecone, ChromaDB
• Cloud & DevOps: AWS, Docker, Kubernetes, CI/CD, Git, GitHub, Linux, MLflow, Weights & Biases
• Methodologies: Agile, Scrum, Unit Testing, System Design, Code Review

EXPERIENCE
Lead Data Scientist | NeuroTech Analytics | 2022 - Present
• Designed and shipped an enterprise RAG pipeline utilizing LangChain, Hugging Face Transformers, and Pinecone vector search, accelerating query resolution by 40%.
• Fine-tuned open-source LLMs using PyTorch and LoRA, achieving 92% domain-specific extraction accuracy.
• Engineered scalable ML pipelines using Scikit-Learn and Pandas on AWS EC2 and S3, processing 10M+ daily transactions.
• Established CI/CD deployment pipelines using Docker and GitHub Actions, slashing release cycles by 50%.

Machine Learning Engineer | DataCore Systems | 2019 - 2022
• Implemented NLP sentiment classification models using BERT and spaCy, boosting marketing campaign conversion by 22%.
• Built customer churn prediction models using Scikit-Learn (Random Forest, XGBoost) and SQL feature stores.
• Deployed RESTful API endpoints with FastAPI and Docker containerization.

EDUCATION
B.S. in Computer Science | University of California, Berkeley
"""

# 2. DevOps Engineer Resume Content
DEVOPS_RESUME = """
JORDAN SMITH
Austin, TX | jordan.smith@cloudops.io | github.com/jordansmith

PROFESSIONAL SUMMARY
Lead Cloud & DevOps Engineer with 6+ years of expertise in Kubernetes, Terraform, AWS cloud architecture, and automated CI/CD pipelines. Specializes in infrastructure as code, zero-downtime microservices deployments, and enterprise monitoring with Prometheus, Grafana, and Datadog.

TECHNICAL SKILLS
• Cloud & Infrastructure: AWS, EC2, S3, RDS, IAM, CloudWatch, EKS, Azure, Google Cloud, Docker, Kubernetes, Helm
• Infrastructure as Code: Terraform, Ansible, CloudFormation
• CI/CD & Automation: GitHub Actions, GitLab CI, Jenkins, ArgoCD, Bash, Python, Linux, Ubuntu
• Observability & Security: Prometheus, Grafana, Datadog, ELK Stack, OpenTelemetry
• Networking & Middleware: Nginx, REST API, Microservices, DNS, SSL/TLS

EXPERIENCE
Senior DevOps Engineer | CloudScale Networks | 2021 - Present
• Architected multi-region Kubernetes clusters on AWS EKS using Terraform, managing 200+ microservices with 99.99% uptime.
• Built end-to-end CI/CD workflows via GitHub Actions and ArgoCD, reducing deployment times from 45 minutes to 4 minutes.
• Automated infrastructure provisioning with Terraform and Ansible, eliminating manual configuration drift.
• Set up comprehensive observability with Prometheus, Grafana, and CloudWatch alerts.

DevOps Systems Engineer | Infrastructure Hub | 2018 - 2021
• Managed AWS EC2, S3, and RDS infrastructure supporting 2M+ active mobile users.
• Containerized legacy monolith into Docker containers and orchestrated deployment using Docker Compose and Kubernetes.
• Implemented automated security compliance scans and vulnerability assessments in Jenkins pipelines.

EDUCATION
B.S. in Information Systems | University of Texas at Austin
"""

# 3. Backend Engineer Resume Content
BACKEND_RESUME = """
PRIYA PATEL
Seattle, WA | priya.patel@devmail.com | github.com/priyapatel

PROFESSIONAL SUMMARY
Backend Software Engineer with 4+ years of experience designing robust RESTful APIs, microservices, and distributed data systems. Expert in Python, FastAPI, Django, PostgreSQL, Redis, and Apache Kafka. Passionate about clean architecture, TDD, and scalable system design.

TECHNICAL SKILLS
• Languages: Python, Go, SQL, Bash
• Frameworks: FastAPI, Django, Flask, Express, Node.js
• Databases: PostgreSQL, MySQL, Redis, MongoDB, Elasticsearch
• Distributed Systems & Messaging: Apache Kafka, RabbitMQ, Celery
• Cloud & Tools: Docker, Docker Compose, AWS, Git, GitHub, Linux, Nginx
• Practices: Unit Testing, TDD, Agile, Scrum, System Design, Microservices, OOP

EXPERIENCE
Senior Backend Developer | FinTech Nexus | 2022 - Present
• Designed distributed payment processing microservices handling $50M+ monthly transaction volume using FastAPI and PostgreSQL.
• Implemented event-driven message queuing with Apache Kafka and Celery, decoupling synchronous dependencies.
• Reduced database query latency by 65% through Redis caching and PostgreSQL indexing strategies.
• Authored comprehensive unit testing suites with pytest, maintaining 95% test coverage under strict TDD practices.

Backend Software Engineer | StreamLine Media | 2020 - 2022
• Developed REST APIs using Django and Django REST Framework for content streaming service.
• Managed containerized environments with Docker and deployed services onto AWS EC2 and RDS.
• Conducted peer code reviews and mentored junior developers on OOP principles and clean code.

EDUCATION
B.S. in Software Engineering | University of Washington
"""

# 4. Job Descriptions
SENIOR_DATA_SCIENTIST_JD = """
Job Title: Senior Data Scientist - AI & NLP
Company: Apex Artificial Intelligence
Location: Remote (US)

About the Role:
We are seeking an experienced Senior Data Scientist to spearhead our Generative AI and NLP initiatives. In this role, you will design and deploy state-of-the-art machine learning models, RAG systems, and transformer-based architectures to power next-generation intelligent applications.

Key Responsibilities:
• Build and fine-tune large language models (LLM) and transformer architectures (BERT, GPT) for custom domain tasks.
• Design and productionize Retrieval-Augmented Generation (RAG) pipelines using LangChain, Hugging Face, and vector databases like Pinecone or ChromaDB.
• Develop end-to-end Machine Learning pipelines from feature engineering to model deployment using PyTorch, Scikit-Learn, and Python.
• Deploy scalable inference microservices using FastAPI and Docker on AWS cloud infrastructure (EC2, S3, EKS).
• Collaborate with cross-functional teams in an Agile / Scrum framework to translate business needs into technical ML solutions.
• Implement robust MLOps practices, model monitoring, and experiment tracking with MLflow or Weights & Biases.

Qualifications & Skills:
• 4+ years of professional experience in Data Science, Machine Learning, and NLP.
• Expert proficiency in Python, SQL, PyTorch, and Scikit-Learn.
• Deep understanding of Deep Learning, Transformers, RAG, and Vector Embeddings.
• Hands-on experience with Docker, Kubernetes, AWS, and CI/CD pipelines.
• Strong foundation in Data Structures, System Design, and Unit Testing.
• Excellent communication and problem-solving abilities.
"""

SENIOR_DEVOPS_JD = """
Job Title: Senior Cloud & DevOps Engineer
Company: CloudVanguard Enterprise
Location: Remote

Role Overview:
CloudVanguard is looking for a Senior DevOps Engineer to scale our multi-cloud infrastructure and streamline developer operations. You will automate cloud environments, orchestrate Kubernetes clusters, and build hardened CI/CD pipelines.

Key Requirements:
• Expert knowledge of AWS cloud services: EC2, S3, RDS, IAM, EKS, and CloudWatch.
• Production experience with Kubernetes (K8s), Helm charts, and Docker containerization.
• Extensive experience writing Infrastructure as Code using Terraform and Ansible.
• Proven capability building automated CI/CD pipelines with GitHub Actions and ArgoCD.
• Strong proficiency in Linux administration, Bash scripting, and Python automation.
• Experience setting up monitoring, metrics, and alerting with Prometheus, Grafana, and Datadog.
• Solid background in Microservices architecture, Nginx, and System Design.
"""

BACKEND_PYTHON_JD = """
Job Title: Senior Backend Engineer (Python / FastAPI)
Company: DataWave Technologies
Location: New York, NY / Hybrid

Job Description:
DataWave is searching for a talented Senior Backend Engineer to architect high-throughput APIs and distributed data pipelines.

Requirements:
• 4+ years of software development experience with Python, FastAPI, and Django.
• Strong experience with relational and NoSQL databases: PostgreSQL, Redis, and MongoDB.
• Deep expertise with message brokers and event streaming: Apache Kafka, RabbitMQ, and Celery.
• Containerization and cloud deployment experience with Docker, AWS, and Git.
• Commitment to Test-Driven Development (TDD), Unit Testing, and Agile methodologies.
• Strong knowledge of OOP, Data Structures, Algorithms, and System Design.
"""


def create_pdf(text: str, filename: str):
    """Generates a clean PDF resume."""
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    width = pdf.epw
    
    for line in text.strip().split("\n"):
        clean_l = line.encode('latin-1', 'replace').decode('latin-1').strip()
        if not clean_l:
            pdf.ln(2)
            continue
        if clean_l.isupper() and len(clean_l) < 40 and "@" not in clean_l:
            pdf.set_font("Helvetica", style="B", size=12)
            pdf.ln(3)
            pdf.cell(width, 7, clean_l, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=10)
        else:
            pdf.multi_cell(width, 5, clean_l)
            
    pdf.output(filename)


def create_docx(text: str, filename: str):
    """Generates a clean DOCX resume."""
    doc = docx.Document()
    for line in text.strip().split("\n"):
        clean_l = line.strip()
        if clean_l.isupper() and len(clean_l) < 40 and not "@" in clean_l:
            p = doc.add_paragraph()
            run = p.add_run(clean_l)
            run.bold = True
            run.font.size = docx.shared.Pt(12)
        else:
            doc.add_paragraph(clean_l)
    doc.save(filename)


def main():
    print("Generating sample resumes and job descriptions...")

    # Write text resumes
    with open(os.path.join(RESUMES_DIR, "data_scientist_resume.txt"), "w", encoding="utf-8") as f:
        f.write(DATA_SCIENTIST_RESUME)
    with open(os.path.join(RESUMES_DIR, "devops_engineer_resume.txt"), "w", encoding="utf-8") as f:
        f.write(DEVOPS_RESUME)
    with open(os.path.join(RESUMES_DIR, "backend_engineer_resume.txt"), "w", encoding="utf-8") as f:
        f.write(BACKEND_RESUME)

    # Generate PDF resumes
    create_pdf(DATA_SCIENTIST_RESUME, os.path.join(RESUMES_DIR, "data_scientist_resume.pdf"))
    create_pdf(DEVOPS_RESUME, os.path.join(RESUMES_DIR, "devops_engineer_resume.pdf"))

    # Generate DOCX resume
    create_docx(BACKEND_RESUME, os.path.join(RESUMES_DIR, "backend_engineer_resume.docx"))

    # Write Job Descriptions
    with open(os.path.join(JDS_DIR, "senior_data_scientist_jd.txt"), "w", encoding="utf-8") as f:
        f.write(SENIOR_DATA_SCIENTIST_JD)
    with open(os.path.join(JDS_DIR, "senior_devops_jd.txt"), "w", encoding="utf-8") as f:
        f.write(SENIOR_DEVOPS_JD)
    with open(os.path.join(JDS_DIR, "backend_python_jd.txt"), "w", encoding="utf-8") as f:
        f.write(BACKEND_PYTHON_JD)

    print("Sample generation complete!")


if __name__ == "__main__":
    main()
